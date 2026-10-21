
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md"}


def extract_text_from_file(file_path: str) -> str:
    """
    Extract readable text from a supported document.

    Supports PDF, TXT, DOCX, and Markdown files.
    Returns extracted text as a string.
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    if not path.is_file():
        raise FileNotFoundError("The uploaded document was not found.")

    if extension == ".pdf":
        reader = PdfReader(str(path))
        extracted_text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif extension == ".docx":
        document = Document(str(path))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_rows = [
            " | ".join(
                cell.text.strip()
                for cell in row.cells
            )
            for table in document.tables
            for row in table.rows
        ]

        extracted_text = "\n".join(paragraphs + table_rows)

    else:
        # TXT and Markdown are both plain-text formats.
        extracted_text = path.read_text(encoding="utf-8-sig")

    extracted_text = extracted_text.strip()

    if not extracted_text:
        raise ValueError(
            "No readable text was found in the document. "
            "The PDF may be scanned or the file may be empty."
        )

    return extracted_text
